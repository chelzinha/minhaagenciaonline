"""Gera a base do Comparador de Tarifas para o D1 (agf-balcao).

Entradas (pasta 02 - Tarifas e Tabelas de Preços/2026 e entrada/):
  Tarifas_Encomendas_Platinum_Clube_2026.xlsx  aba BASE 2026  (16 pacotes: SEDEX, PAC, Mini Envios)
  Tarifas_Enc_APP_2026.xlsx                    aba BASE-2026  (SEDEX APP e PAC APP)
  entrada/ (ou ../01_ENTRADA no Drive)
  classificacao_cidades_2023.csv               classe A+, A, B, C por município
  entrada/faixa_por_uf.csv                     faixa N/P (corredor Fortaleza E+) e faixa I (Matriz CE)
  entrada/regras.csv                           ad valorem, indenização, divisores, Mini Envios

Saídas:
  saida/02_D1_CSV/*.csv
  saida/0005_comparador_seed_<ano>.sql   (copiar para cloudflare/balcao-api/migrations)

Uso:
  pip install openpyxl
  python gerar_comparador_d1.py --contrato "<xlsx contrato>" --app "<xlsx app>" --vigencia 2026-04-12
"""
import argparse, csv, datetime, os, re, unicodedata
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
PACOTES = ['CLUBE CORREIOS', 'PLATINUM', 'DIAMANTE START 1', 'DIAMANTE START 2', 'DIAMANTE 1', 'DIAMANTE 2',
           'DIAMANTE 3', 'DIAMANTE 4', 'INFINITE 1', 'INFINITE 2', 'INFINITE 3', 'INFINITE 4', 'INFINITE 5',
           'INFINITE 6', 'INFINITE 7', 'INFINITE 8']
BLOCOS_CONTRATO = {'SEDEX': 'SEDEX', 'PAC': 'PAC', 'MINI ENVIOS': 'MINI'}
BLOCOS_APP = {'SEDEX APP': 'SEDEX', 'PAC APP': 'PAC'}
COLUNA_RE = re.compile(r'[LENPI][1-4]')


def norm(s):
    s = ''.join(c for c in unicodedata.normalize('NFKD', str(s or '')) if not unicodedata.combining(c))
    return ' '.join(re.sub(r'[^A-Z0-9]+', ' ', s.upper()).split())


def codigo_tabela(pacote):
    return 'APP' if pacote == 'APP' else norm(pacote).replace(' ', '_')


def num(v):
    if v is None or v == '':
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    return round(float(str(v).strip().replace('.', '').replace(',', '.')), 2)


def faixa(texto):
    """'0 a 300' -> ('FAIXA', 300); '9.001 A 10.000' -> ('FAIXA', 10000); 'Kg Adicional' -> ('KG_ADICIONAL', 10000)."""
    t = str(texto).strip().lower()
    if 'adicional' in t:
        return 'KG_ADICIONAL', 10000
    nums = re.findall(r'\d+', t.replace('.', ''))
    return ('FAIXA', int(nums[-1])) if nums else (None, None)


def ler_blocos(ws, mapa):
    """Lê uma aba BASE: linha 1 = nome do bloco, linha 2 = colunas L1..I4, linhas 3+ = chave, pacote, faixa, preços."""
    linhas = list(ws.iter_rows(values_only=True))
    r0, r1 = linhas[0], linhas[1]
    inicios = [(j, str(v).strip()) for j, v in enumerate(r0) if v and j > 0]
    saida = []
    for i, (c, nome) in enumerate(inicios):
        if nome not in mapa:
            continue
        fim = inicios[i + 1][0] if i + 1 < len(inicios) else len(r0)
        cols = [(j, str(r1[j]).strip()) for j in range(c, fim) if r1[j] and COLUNA_RE.fullmatch(str(r1[j]).strip())]
        for r in linhas[2:]:
            pac, fx = (r[c + 1] if c + 1 < len(r) else None), (r[c + 2] if c + 2 < len(r) else None)
            if not pac or fx is None:
                continue
            tipo, peso = faixa(fx)
            if not tipo:
                continue
            for j, col in cols:
                v = num(r[j])
                if v is not None:
                    saida.append((str(pac).strip(), mapa[nome], col, tipo, peso, v))
    return saida


ENTRADA = next((d for d in [os.path.join(AQUI, 'entrada'), os.path.join(AQUI, '..', '01_ENTRADA')] if os.path.isdir(d)), os.path.join(AQUI, 'entrada'))


def ler_csv(nome):
    with open(os.path.join(ENTRADA, nome), encoding='utf-8') as f:
        return list(csv.DictReader(f, delimiter=';'))


def q(v):
    if v is None or v == '':
        return 'NULL'
    if isinstance(v, (int, float)):
        return repr(v) if isinstance(v, int) else f'{v:.4f}'.rstrip('0').rstrip('.')
    return "'" + str(v).replace("'", "''") + "'"


def inserts(tabela, colunas, linhas, lote=150):
    out = []
    for i in range(0, len(linhas), lote):
        vals = ',\n'.join('(' + ','.join(q(v) for v in l) + ')' for l in linhas[i:i + lote])
        out.append(f'INSERT INTO {tabela} ({",".join(colunas)}) VALUES\n{vals};')
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--contrato', required=True)
    ap.add_argument('--app', required=True)
    ap.add_argument('--vigencia', default='2026-04-12')
    a = ap.parse_args()

    contrato = ler_blocos(openpyxl.load_workbook(a.contrato, read_only=True, data_only=True)['BASE 2026'], BLOCOS_CONTRATO)
    app = ler_blocos(openpyxl.load_workbook(a.app, read_only=True, data_only=True)['BASE-2026'], BLOCOS_APP)

    pacotes_lidos = sorted({r[0] for r in contrato})
    faltando = [p for p in PACOTES if p not in pacotes_lidos]
    assert not faltando, f'Pacotes ausentes na planilha: {faltando}'
    assert len(app) > 300, 'Tabela App incompleta'

    tarifas = [(codigo_tabela(p), s, c, t, g, v) for p, s, c, t, g, v in contrato]
    tarifas += [('APP', s, c, t, g, v) for _, s, c, t, g, v in app]
    chaves = {(r[0], r[1], r[2], r[3], r[4]) for r in tarifas}
    assert len(chaves) == len(tarifas), 'Linha de tarifa duplicada'

    tabelas = [(codigo_tabela(p), p.title().replace('Start', 'Start'), 'CONTRATO', i + 1,
                'Mesmos preços do Platinum na tabela 2026' if p == 'CLUBE CORREIOS' else '', a.vigencia)
               for i, p in enumerate(PACOTES)]
    tabelas.append(('APP', 'Correios App', 'APP', 99, 'Pré-postagem à vista pelo App. Cobra sempre o peso cúbico (divisor 7000).', a.vigencia))

    classes = [(r['uf'], norm(r['municipio']), r['municipio'], r['classe'], r['fonte']) for r in ler_csv('classificacao_cidades_2023.csv')]
    faixas = [(r['uf'], int(r['faixa_capital']), int(r['faixa_interior']), r['observacao'])
              for r in ler_csv('faixa_por_uf.csv') if r['faixa_capital']]
    regras = [(r['chave'], float(r['valor']), r['descricao'], r['fonte']) for r in ler_csv('regras.csv')]

    saida = os.path.join(AQUI, 'saida')
    os.makedirs(os.path.join(saida, '02_D1_CSV'), exist_ok=True)
    for nome, cab, linhas in [
        ('01_tabelas.csv', ['codigo', 'nome', 'tipo', 'ordem', 'observacao', 'vigencia'], tabelas),
        ('02_tarifas.csv', ['tabela', 'servico', 'coluna', 'tipo', 'peso_max_g', 'preco'], tarifas),
        ('03_faixa_por_uf.csv', ['uf', 'faixa_capital', 'faixa_interior', 'observacao'], faixas),
        ('04_classe_cidade.csv', ['uf', 'municipio_norm', 'municipio', 'classe', 'fonte'], classes),
        ('05_regras.csv', ['chave', 'valor', 'descricao', 'fonte'], regras)]:
        with open(os.path.join(saida, '02_D1_CSV', nome), 'w', newline='', encoding='utf-8') as f:
            w = csv.writer(f, delimiter=';'); w.writerow(cab); w.writerows(linhas)

    ano = a.vigencia[:4]
    sql = [f'-- Comparador de Tarifas - base {a.vigencia} | gerado em {datetime.date.today().isoformat()} por gerar_comparador_d1.py',
           f'-- {len(tarifas)} tarifas ({len(PACOTES)} pacotes de contrato + App), {len(classes)} cidades classificadas',
           'DELETE FROM cmp_tarifas;', 'DELETE FROM cmp_tabelas;', 'DELETE FROM cmp_faixa_uf;',
           'DELETE FROM cmp_classe_cidade;', 'DELETE FROM cmp_regras;']
    sql += inserts('cmp_tabelas', ['codigo', 'nome', 'tipo', 'ordem', 'observacao', 'vigencia'], tabelas)
    sql += inserts('cmp_tarifas', ['tabela', 'servico', 'coluna', 'tipo', 'peso_max_g', 'preco'], tarifas)
    sql += inserts('cmp_faixa_uf', ['uf', 'faixa_capital', 'faixa_interior', 'observacao'], faixas)
    sql += inserts('cmp_classe_cidade', ['uf', 'municipio_norm', 'municipio', 'classe', 'fonte'], classes)
    sql += inserts('cmp_regras', ['chave', 'valor', 'descricao', 'fonte'], regras)
    sql.append(f"INSERT INTO cmp_carga_log (vigencia, linhas_tarifas, linhas_cidades, observacao) VALUES "
               f"({q(a.vigencia)}, {len(tarifas)}, {len(classes)}, {q(os.path.basename(a.contrato) + ' + ' + os.path.basename(a.app))});")
    destino = os.path.join(saida, f'0005_comparador_seed_{ano}.sql')
    with open(destino, 'w', encoding='utf-8') as f:
        f.write('\n'.join(sql) + '\n')
    print(f'OK: {len(tarifas)} tarifas, {len(tabelas)} tabelas, {len(classes)} cidades, {len(regras)} regras -> {destino}')


if __name__ == '__main__':
    main()
